"""Turn an uploaded PDF / DOCX / TXT into a ResumeProfile, in memory only.

Owner: Charith Reddy Mummadireddy Reddivari.
Layout signals (columns, tables, text boxes, header/footer) are gathered because those
are the commonly cited ATS parsing hazards (design doc, "Résumé Ingestion").
"""
from __future__ import annotations

import io
from pathlib import Path

from ..models import ResumeProfile
from .redact import detect_contacts, redact
from .sectioning import split_sections

MIN_TEXT_CHARS = 150
SUPPORTED = {".pdf", ".docx", ".txt"}


class ResumeReadError(Exception):
    """Raised with a user-safe message when a file can't be used."""


def load_resume(filename: str, data: bytes) -> ResumeProfile:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ResumeReadError("Unsupported file type. Please upload a PDF, DOCX, or TXT file.")
    if not data:
        raise ResumeReadError("The file is empty.")
    try:
        if ext == ".pdf":
            raw, layout = _read_pdf(data)
        elif ext == ".docx":
            raw, layout = _read_docx(data)
        else:
            raw, layout = _read_txt(data)
    except ResumeReadError:
        raise
    except Exception as e:  # corrupt or password-protected files
        raise ResumeReadError(f"Couldn't read that file ({type(e).__name__}). Is it corrupt or password-protected?") from e

    layout["image_only"] = len(raw.strip()) < MIN_TEXT_CHARS and ext == ".pdf"
    if len(raw.strip()) < MIN_TEXT_CHARS and ext == ".pdf":
        raise ResumeReadError(
            "This PDF has almost no extractable text - it looks like a scanned image. "
            "ATS software can't read image-only files either. Export a text-based PDF or upload DOCX/TXT."
        )
    if len(raw.strip()) < 20:
        raise ResumeReadError("I couldn't find any text in that file.")

    contacts = detect_contacts(raw)
    return ResumeProfile(
        file_type=ext.lstrip("."),
        raw_text=raw,
        redacted_text=redact(raw, contacts),
        contacts=contacts,
        sections=split_sections(raw),
        layout=layout,
    )


def _blank_layout(file_type: str) -> dict:
    return {"file_type": file_type, "pages": None, "multi_column": False, "tables": 0,
            "textboxes": 0, "header_footer_text": "", "image_only": False}


def _read_txt(data: bytes) -> tuple[str, dict]:
    for enc in ("utf-8", "latin-1"):
        try:
            return data.decode(enc), _blank_layout("txt")
        except UnicodeDecodeError:
            continue
    raise ResumeReadError("Couldn't decode the text file.")


def _read_docx(data: bytes) -> tuple[str, dict]:
    import docx

    d = docx.Document(io.BytesIO(data))
    layout = _blank_layout("docx")
    paras = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        bullet = "List" in (p.style.name if p.style is not None else "")
        paras.append(("• " + text) if bullet and not text.startswith(("•", "-", "*")) else text)
    layout["tables"] = len(d.tables)
    table_text = []
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                table_text.append(" ".join(cells))
    body_xml = d.element.body.xml
    layout["textboxes"] = body_xml.count("<w:txbxContent")
    layout["multi_column"] = any(
        (s._sectPr.xpath("./w:cols/@w:num") or ["1"])[0] not in ("1", "") for s in d.sections
    )
    hf = []
    for s in d.sections:
        for part in (s.header, s.footer):
            hf.extend(p.text.strip() for p in part.paragraphs if p.text.strip())
    layout["header_footer_text"] = "\n".join(hf)
    # Table text is appended: parsers often read tables out of order, which is what the ATS rule flags.
    return "\n".join(paras + table_text), layout


def _read_pdf(data: bytes) -> tuple[str, dict]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ResumeReadError("This PDF is password-protected.")
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    layout = _blank_layout("pdf")
    layout["pages"] = len(reader.pages)
    layout["multi_column"] = _pdf_multi_column(data)
    return text, layout


def _pdf_multi_column(data: bytes) -> bool:
    """Heuristic: many text lines have a wide horizontal gap near the page centre."""
    try:
        import pdfplumber
    except ImportError:  # pragma: no cover
        return False
    wide_lines = total_lines = 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:2]:
                words = page.extract_words()
                width = float(page.width)
                rows: dict[int, list[dict]] = {}
                for w in words:
                    rows.setdefault(round(w["top"] / 3), []).append(w)
                for row in rows.values():
                    row.sort(key=lambda w: w["x0"])
                    if len(row) < 4:
                        continue
                    total_lines += 1
                    for a, b in zip(row, row[1:]):
                        gap = b["x0"] - a["x1"]
                        mid = (a["x1"] + b["x0"]) / 2 / width
                        if gap > 0.06 * width and 0.3 < mid < 0.7:
                            wide_lines += 1
                            break
    except Exception:
        return False
    return total_lines >= 6 and wide_lines / total_lines >= 0.4
