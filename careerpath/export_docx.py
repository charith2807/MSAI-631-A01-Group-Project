"""Export an accepted-changes résumé as a single-column, ATS-friendly DOCX (no tables/text boxes).

Owner: Charith Reddy Mummadireddy Reddivari.
Contact details are re-attached locally here - they were never sent to the model.
"""
from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt

from .models import ResumeProfile, TailoredResume

TITLES = {"summary": "Summary", "experience": "Experience", "education": "Education", "skills": "Skills",
          "projects": "Projects", "certifications": "Certifications"}


def apply_accepted(resume: ResumeProfile, tailored: TailoredResume | None, accepted_ids: set[str]) -> dict[str, str]:
    """Map of line id -> replacement text, only for accepted, non-rejected changes."""
    out: dict[str, str] = {}
    if tailored:
        for ch in tailored.changes:
            if ch.id in accepted_ids and ch.status != "rejected":
                out[ch.id] = ch.proposed
    return out


def build_docx(resume: ResumeProfile, tailored: TailoredResume | None = None, accepted_ids: set[str] | None = None) -> bytes:
    repl = apply_accepted(resume, tailored, accepted_ids or set())
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for sec in resume.sections:
        if sec.key == "header":
            for i, ln in enumerate(sec.lines):
                p = doc.add_paragraph()
                run = p.add_run(ln.text)
                if i == 0:
                    run.bold = True
                    run.font.size = Pt(16)
            continue
        doc.add_heading(TITLES.get(sec.key, sec.heading.title()), level=2)
        if sec.key == "summary" and "summary" in repl:
            doc.add_paragraph(repl["summary"])
            continue
        for ln in sec.lines:
            text = repl.get(ln.id, ln.text)
            if ln.kind == "bullet":
                doc.add_paragraph(text, style="List Bullet")
            else:
                doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
