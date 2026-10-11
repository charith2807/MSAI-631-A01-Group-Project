"""Regenerate the synthetic layout fixtures in tests/fixtures/layouts/ (all fictional people).

Usage: python scripts/make_fixtures.py     (needs requirements-dev.txt)
Owner: Jangmu Sherpa (fixtures) with code support from Charith Reddy Mummadireddy Reddivari.
"""
from pathlib import Path

import docx
from docx.shared import Pt
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "layouts"
OUT.mkdir(parents=True, exist_ok=True)

NAME = "Maya Chen"
CONTACT = "maya.chen@example.com | (214) 555-0101 | linkedin.com/in/maya-chen-example"
SUMMARY = "Data analyst with three years of experience turning retail sales data into clear reports for store managers."
JOBS = [
    ("Data Analyst, Northfield Retail, Jan 2023 - Present", [
        "Built SQL dashboards in Tableau that cut weekly reporting time by 40%",
        "Cleaned and merged sales data from 12 stores using Python and Pandas",
        "Presented data analysis findings to regional managers and improved communication with head office",
        "Maintained Excel forecasting models used by a team of 8 buyers"]),
    ("Junior Analyst, Brightside Foods, Jun 2021 - Dec 2022", [
        "Reduced report errors by 25% by validating inputs with SQL checks",
        "Created weekly sales summaries for 30 product lines",
        "Documented 14 recurring reports and trained 3 new analysts"]),
]
EDU = "B.S. Statistics, Example State University, 2021"
SKILLS = "SQL, Python, Pandas, Excel, Tableau, Agile"


def plain_lines(headings=("SUMMARY", "EXPERIENCE", "EDUCATION", "SKILLS"), mixed_dates=False, bullets=True):
    s, e, ed, sk = headings
    out = [NAME, CONTACT, "", s, SUMMARY, "", e]
    for i, (title, bl) in enumerate(JOBS):
        if mixed_dates and i == 1:
            title = "Junior Analyst, Brightside Foods, 06/2021 - December 2022"
        out.append(title)
        out += [("- " + b) if bullets else b + "." for b in bl]
    out += ["", ed, EDU, "", sk, SKILLS]
    return out


def write_txt(name, lines):
    (OUT / name).write_text("\n".join(lines) + "\n", encoding="utf-8")


def single_col_pdf(path):
    c = canvas.Canvas(str(path), pagesize=letter)
    y = 750
    for ln in plain_lines():
        c.setFont("Helvetica-Bold" if ln.isupper() and ln else "Helvetica", 11)
        c.drawString(60, y, ("• " + ln[2:]) if ln.startswith("- ") else ln)
        y -= 16
    c.save()


def two_col_pdf(path):
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica-Bold", 16); c.drawString(60, 750, NAME)
    c.setFont("Helvetica", 10); c.drawString(60, 735, CONTACT)
    left = ["EXPERIENCE", JOBS[0][0][:55], *["- " + b[:52] for b in JOBS[0][1]], JOBS[1][0][:55], *["- " + b[:52] for b in JOBS[1][1]]]
    right = ["SKILLS", "SQL, Python, Pandas", "Excel, Tableau, Agile", "", "EDUCATION", "B.S. Statistics", "Example State University", "2021",
             "", "SUMMARY", "Data analyst with three years", "of experience in retail", "reporting for managers"]
    y = 700
    for ln in left:
        c.setFont("Helvetica", 9); c.drawString(60, y, ln); y -= 16
    y = 700
    for ln in right:
        c.setFont("Helvetica", 9); c.drawString(340, y, ln); y -= 16
    c.save()


def image_only_pdf(path):
    c = canvas.Canvas(str(path), pagesize=letter)
    c.rect(60, 400, 400, 300, fill=1)  # no text at all, stands in for a scanned page
    c.save()


def clean_docx(path):
    d = docx.Document()
    d.styles["Normal"].font.size = Pt(11)
    for ln in plain_lines():
        if ln.startswith("- "):
            d.add_paragraph(ln[2:], style="List Bullet")
        elif ln.isupper() and ln:
            d.add_heading(ln.title(), level=2)
        elif ln:
            d.add_paragraph(ln)
    d.save(path)


def table_docx(path):
    d = docx.Document()
    d.add_paragraph(NAME); d.add_paragraph(CONTACT)
    d.add_heading("Summary", level=2); d.add_paragraph(SUMMARY)
    d.add_heading("Experience", level=2)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = JOBS[0][0]
    t.rows[0].cells[1].text = "\n".join("- " + b for b in JOBS[0][1])
    for title, bl in JOBS[1:]:
        d.add_paragraph(title)
        for b in bl:
            d.add_paragraph(b, style="List Bullet")
    d.add_heading("Education", level=2); d.add_paragraph(EDU)
    d.add_heading("Skills", level=2); d.add_paragraph(SKILLS)
    d.save(path)


def header_contact_docx(path):
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = CONTACT
    for ln in plain_lines():
        if ln == CONTACT:
            continue
        if ln.startswith("- "):
            d.add_paragraph(ln[2:], style="List Bullet")
        elif ln.isupper() and ln:
            d.add_heading(ln.title(), level=2)
        elif ln:
            d.add_paragraph(ln)
    d.save(path)


if __name__ == "__main__":
    write_txt("clean.txt", plain_lines())
    single_col_pdf(OUT / "clean.pdf")
    clean_docx(OUT / "clean.docx")
    two_col_pdf(OUT / "two_column.pdf")
    image_only_pdf(OUT / "image_only.pdf")
    table_docx(OUT / "table_layout.docx")
    header_contact_docx(OUT / "header_contact.docx")
    write_txt("creative_headings.txt", plain_lines(("MY STORY", "WHERE I'VE WORKED", "SCHOOLING", "TOOLBOX")))
    write_txt("mixed_dates.txt", plain_lines(mixed_dates=True))
    write_txt("no_bullets.txt", plain_lines(bullets=False))
    write_txt("too_short.txt", [NAME, CONTACT, "", "EXPERIENCE", "Data Analyst, Northfield Retail, Jan 2023 - Present",
                                "- Built SQL dashboards in Tableau that cut reporting time by 40%", "", "EDUCATION", EDU, "", "SKILLS", SKILLS])
    print("wrote fixtures to", OUT)
